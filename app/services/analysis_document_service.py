# app/services/analysis_document_service.py

import json
import uuid
import openpyxl
import pdfplumber
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import List, Dict, Any
from docx import Document
from openai import OpenAI
from app.config import settings
from app.models.analysis_document_schema import (
    AnalysisDocumentRequest,
    AnalysisDocumentResponse
)
from app.utils.file_validator import FileValidator

client = OpenAI(api_key=settings.OPENAI_API_KEY)


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ============================================================
# DocumentAnalyzerV3 (AI.ipynb 원본 그대로)
# ============================================================

class DocumentAnalyzerV3:
    """수사 문서 분석 시스템"""

    FORENSIC_PROMPT = """당신은 수사 보조 분석 담당자입니다.

아래에는 사건 개요와 문서 원문이 제공됩니다.

문서에 기재된 모든 객관적 정보를 빠짐없이 추출하여 설명문을 작성하되,
반드시 다음을 포함하십시오:

1. 문서의 종류 및 작성 목적
2. 문서에 기재된 주요 사실 (모든 객관적 정보 포함)
   - 인적 사항 (이름, 생년월일, 주소, 연락처, 직업 등)
   - 날짜/시간 정보 (사건 발생 일시, 작성 일시, 기타 시간 정보)
   - 장소 정보 (주소, 위치, 장소 설명)
   - 물품/증거물 정보 (종류, 수량, 상태, 특징)
   - 상해/손상 정보 (부위, 정도, 치료 내용, 전치 기간)
   - 금액/거래 정보 (가격, 금액, 거래 내역)
   - 차량/물건 식별 정보 (번호판, 시리얼, IMEI 등)
   - 기타 측정 가능한 수치 (크기, 무게, 거리 등)
3. 사건 개요와 직접적으로 관련되는 부분

절대 규칙:
- 문서에 명시된 모든 구체적 정보를 빠짐없이 포함할 것
- 숫자, 날짜, 이름, 주소 등은 정확히 그대로 기록
- 문서에 없는 사실 추가 금지
- 추정, 해석, 판단, 법적 평가 금지
- 사건 개요에 없는 내용을 새로 만들지 말 것
- 목록/번호 사용 금지
- 설명문으로 작성 (6~10문장)

[사건 정보]
{case_info}

[문서 원문]
{document_text}
"""

    def __init__(self, model: str = "gpt-4o-mini", temperature: float = 0.0, chunk_size: int = 3000):
        self.model = model
        self.temperature = temperature
        self.chunk_size = chunk_size

    def extract(self, path: Path) -> str:
        """문서에서 텍스트 추출"""
        ext = path.suffix.lower()

        # PDF
        if ext == ".pdf":
            texts = []
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    t = page.extract_text()
                    if t:
                        texts.append(f"[Page {i}]\n{t}")
            return "\n\n".join(texts)

        # DOCX
        if ext == ".docx":
            doc = Document(path)
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(lines)

        # XLSX/XLS
        if ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(path, data_only=True)
            blocks = []
            for ws in wb.worksheets:
                blocks.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows():
                    values = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        # Excel 날짜 변환
                        if isinstance(cell.value, (int, float)) and 30000 < cell.value < 60000:
                            base = datetime(1899, 12, 30)
                            dt = base + timedelta(days=int(cell.value))
                            values.append(dt.strftime("%Y-%m-%d"))
                        else:
                            values.append(str(cell.value))
                    if values:
                        blocks.append(" | ".join(values))
            return "\n".join(blocks)

        # TXT/CSV/LOG
        if ext in [".txt", ".csv", ".log"]:
            for enc in ["utf-8", "cp949", "euc-kr", "latin1"]:
                try:
                    return path.read_text(encoding=enc)
                except Exception:
                    continue
            raise ValueError("텍스트 인코딩 판독 실패")

        raise ValueError(f"지원하지 않는 문서 형식: {ext}")

    def split_chunks(self, text: str) -> List[str]:
        """텍스트를 청크로 분할 (10% 오버랩)"""
        size = self.chunk_size
        overlap = int(size * 0.1)
        chunks = []
        start = 0
        while start < len(text):
            end = start + size
            chunks.append(text[start:end])
            start = end - overlap
        return chunks

    def _format_case_info(self, case_json: dict) -> str:
        """CASE JSON을 사건 정보 문자열로 변환"""
        if not case_json:
            return "사건 정보 없음"
        
        info_parts = []
        
        if "title" in case_json:
            info_parts.append(f"사건명: {case_json['title']}")
        
        if "type" in case_json:
            info_parts.append(f"사건 유형: {case_json['type']}")
        
        if "description" in case_json:
            info_parts.append(f"사건 개요: {case_json['description']}")
        
        occurred_time = case_json.get('occurred_at', case_json.get('occured_at', ''))
        if occurred_time:
            info_parts.append(f"발생 일시: {occurred_time}")
        
        if "location" in case_json:
            info_parts.append(f"발생 장소: {case_json['location']}")
        
        # summary 추가
        if "summary" in case_json and case_json["summary"]:
            summary = case_json["summary"]
            if isinstance(summary, dict) and any(summary.values()):
                info_parts.append("\n사건 요약:")
                for key, value in summary.items():
                    if value:
                        info_parts.append(f"  - {key}: {value}")
        
        return "\n".join(info_parts)

    def analyze_one(self, text: str, case_json: dict = None) -> str:
        """문서 분석 실행"""
        case_info = self._format_case_info(case_json)
        chunks = self.split_chunks(text)
        partials = []

        for chunk in chunks:
            prompt = self.FORENSIC_PROMPT.format(
                case_info=case_info,
                document_text=chunk
            )

            response = client.chat.completions.create(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                temperature=self.temperature
            )

            partials.append(response.choices[0].message.content.strip())

        return "\n".join(partials)


# ============================================================
# AnalysisDocumentService
# ============================================================

class AnalysisDocumentService:
    @staticmethod
    def analyze_document(request: AnalysisDocumentRequest) -> AnalysisDocumentResponse:
        """
        문서 분석 (AI.ipynb의 analyze_document_to_fixed_json 로직)
        """
        file_path = Path(request.file_path)
        
        # ✅ 파일 검증 추가
        FileValidator.validate_document(request.file_type, request.file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"파일을 찾을 수 없습니다: {request.file_path}")
        
        print(f"\n🔄 문서 분석 시작: {request.origin_name}")
        
        # DocumentAnalyzerV3 초기화
        analyzer = DocumentAnalyzerV3(
            model="gpt-4o-mini",
            temperature=0.0,
            chunk_size=3000
        )
        
        # ✅ analysis_id는 자동 생성
        analysis_id = str(uuid.uuid4())
        created_at = utc_now_iso()
        status = "PROCESSING"
        
        try:
            # 텍스트 추출
            text = analyzer.extract(file_path)
            print(f"✅ 텍스트 추출 완료 ({len(text):,}자)")
            
            # case_json 구성
            case_json = {
                "case_id": request.case_id
            }
            
            # 분석 실행
            analysis_text = analyzer.analyze_one(text, case_json)
            status = "COMPLETED"
            print(f"✅ 분석 완료")
            
        except Exception as e:
            analysis_text = f"분석 실패: {str(e)}"
            status = "PENDING"
            print(f"⚠️ 분석 실패: {e}")
        
        # ✅ AI.ipynb와 동일한 응답 형식
        return AnalysisDocumentResponse(
            analysis_id=analysis_id,
            file_id=request.file_id,
            analysis_type="DOC",
            status=status,
            result_data=analysis_text,
            created_at=created_at
        )